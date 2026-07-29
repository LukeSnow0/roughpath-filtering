"""Load InTaSet Excel trajectories for filtering experiments.

InTaSet does **not** ship i.i.d. Monte Carlo replicates of one experiment.
For averaging we treat non-overlapping time windows of the long 4T training
records as pseudo-realizations (same plant / valve mode, different excitation
segments).
"""

from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path
from typing import Iterable

import numpy as np

try:
    import openpyxl
except ImportError as exc:  # pragma: no cover
    raise ImportError(
        "openpyxl is required to read InTaSet .xlsx files "
        "(uv sync --extra experiments)"
    ) from exc

HERE = Path(__file__).resolve().parent
DATA_ROOT = HERE / "data" / "InTaSet"

# Canonical 4-tank ID paths used by the Pareto experiment.
PATHS_4T = {
    "4T-M1-train": DATA_ROOT
    / "InTaSet-ID"
    / "InTaSet-ID-4T-M1"
    / "InTaSet-ID-4T-M1_Training"
    / "InTaSet-ID-4T-M1_Training.xlsx",
    "4T-M1-test": DATA_ROOT
    / "InTaSet-ID"
    / "InTaSet-ID-4T-M1"
    / "InTaSet-ID-4T-M1_Testing"
    / "InTaSet-ID-4T-M1_Testing.xlsx",
    "4T-M2-train": DATA_ROOT
    / "InTaSet-ID"
    / "InTaSet-ID-4T-M2"
    / "InTaSet-ID-4T-M2_Training"
    / "InTaSet-ID-4T-M2_Training.xlsx",
    "4T-M2-test": DATA_ROOT
    / "InTaSet-ID"
    / "InTaSet-ID-4T-M2"
    / "InTaSet-ID-4T-M2_Testing"
    / "InTaSet-ID-4T-M2_Testing.xlsx",
}


@dataclass(frozen=True)
class IntaSetTrajectory:
    """One continuous recording (or a window of one).

    Attributes
    ----------
    t_sec : (T,) time in seconds, starting at 0 for windows
    heights : (T, 4) tank levels [cm] (sensor readings; may be slightly negative)
    pumps : (T, 2) pump commands in [0, 100]
    source : provenance string, e.g. ``4T-M2-train#w0``
    """

    t_sec: np.ndarray
    heights: np.ndarray
    pumps: np.ndarray
    source: str

    @property
    def n(self) -> int:
        return int(self.t_sec.shape[0])

    @property
    def duration_sec(self) -> float:
        return float(self.t_sec[-1] - self.t_sec[0])

    @property
    def dt_sec(self) -> float:
        return float(np.median(np.diff(self.t_sec)))


def _read_xlsx_numeric(path: Path) -> tuple[list[str], np.ndarray]:
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    rows = list(ws.iter_rows(values_only=True))
    wb.close()
    if not rows:
        raise ValueError(f"empty workbook: {path}")
    header = [("" if c is None else str(c).strip()) for c in rows[0]]
    data_rows = []
    for r in rows[1:]:
        if r is None or all(v is None or v == "" for v in r):
            continue
        vals = []
        for v in r[: len(header)]:
            if v is None or v == "":
                vals.append(np.nan)
            else:
                vals.append(float(v))
        data_rows.append(vals)
    return header, np.asarray(data_rows, dtype=np.float64)


def _col_matches(name: str, *needles: str) -> bool:
    low = name.lower().replace("_", " ").replace("-", " ")
    return any(n in low for n in needles)


def load_trajectory(path: Path | str, *, source: str | None = None) -> IntaSetTrajectory:
    """Load a 4-tank InTaSet sheet into ``(t, heights, pumps)``."""
    path = Path(path)
    header, arr = _read_xlsx_numeric(path)
    if arr.ndim != 2 or arr.shape[1] < 6:
        raise ValueError(f"expected ≥6 columns in {path}, got shape {arr.shape}")

    # Time may have a blank header (4T-M1 training).
    time_idx = next((i for i, h in enumerate(header) if _col_matches(h, "time")), None)
    if time_idx is None:
        # Blank first column that looks like a uniform grid.
        if header[0] == "" and arr.shape[0] > 2:
            d0 = np.diff(arr[:, 0])
            if np.nanmedian(d0) > 0:
                time_idx = 0
    tank_idxs = [i for i, h in enumerate(header) if _col_matches(h, "tank")]
    pump_idxs = [i for i, h in enumerate(header) if _col_matches(h, "pump")]

    if len(tank_idxs) < 4 or len(pump_idxs) < 2:
        # Fall back to positional layout: [time?, h1..h4, u1, u2]
        if time_idx == 0 or (time_idx is None and arr.shape[1] >= 7):
            time_idx = 0 if time_idx is None else time_idx
            tank_idxs = list(range(1, 5))
            pump_idxs = list(range(5, 7))
        elif arr.shape[1] >= 6:
            time_idx = None
            tank_idxs = list(range(0, 4))
            pump_idxs = list(range(4, 6))
        else:
            raise ValueError(f"cannot locate tank/pump columns in {path}: {header}")

    if time_idx is None:
        t = np.arange(arr.shape[0], dtype=np.float64) * 0.1
    else:
        t = arr[:, time_idx].astype(np.float64)
    heights = arr[:, tank_idxs[:4]].astype(np.float64)
    pumps = arr[:, pump_idxs[:2]].astype(np.float64)

    # Drop rows with NaNs in core signals.
    ok = np.isfinite(t) & np.all(np.isfinite(heights), axis=1) & np.all(
        np.isfinite(pumps), axis=1
    )
    t, heights, pumps = t[ok], heights[ok], pumps[ok]
    if t.size < 2:
        raise ValueError(f"too few valid rows in {path}")

    # Enforce strictly increasing time (drop duplicates / reversals).
    order = np.argsort(t)
    t, heights, pumps = t[order], heights[order], pumps[order]
    keep = np.ones(t.shape[0], dtype=bool)
    keep[1:] = np.diff(t) > 1e-12
    t, heights, pumps = t[keep], heights[keep], pumps[keep]

    return IntaSetTrajectory(
        t_sec=t,
        heights=heights,
        pumps=pumps,
        source=source or path.name,
    )


def load_named(name: str) -> IntaSetTrajectory:
    if name not in PATHS_4T:
        raise KeyError(f"unknown trajectory {name!r}; choose from {sorted(PATHS_4T)}")
    return load_trajectory(PATHS_4T[name], source=name)


def window_trajectory(
    traj: IntaSetTrajectory,
    *,
    start_sec: float,
    duration_sec: float,
    window_id: int,
) -> IntaSetTrajectory:
    """Extract ``[start, start+duration]`` and shift time to start at 0."""
    t0 = float(traj.t_sec[0]) + start_sec
    t1 = t0 + duration_sec
    m = (traj.t_sec >= t0 - 1e-9) & (traj.t_sec <= t1 + 1e-9)
    if np.count_nonzero(m) < 2:
        raise ValueError(
            f"window [{start_sec}, {start_sec + duration_sec}] empty for {traj.source}"
        )
    t = traj.t_sec[m] - traj.t_sec[m][0]
    return IntaSetTrajectory(
        t_sec=t,
        heights=traj.heights[m].copy(),
        pumps=traj.pumps[m].copy(),
        source=f"{traj.source}#w{window_id}",
    )


def iter_windows(
    traj: IntaSetTrajectory,
    *,
    duration_sec: float,
    n_windows: int | None = None,
    stride_sec: float | None = None,
) -> list[IntaSetTrajectory]:
    """Non-overlapping (by default) windows covering the recording."""
    stride = duration_sec if stride_sec is None else float(stride_sec)
    total = traj.duration_sec
    max_n = int(total // duration_sec) if stride == duration_sec else None
    if n_windows is None:
        if max_n is None:
            raise ValueError("n_windows required when stride != duration")
        n_windows = max_n
    if n_windows < 1:
        raise ValueError("n_windows must be ≥ 1")
    if max_n is not None and n_windows > max_n:
        raise ValueError(
            f"requested {n_windows} windows of {duration_sec}s but "
            f"{traj.source} only supports {max_n} non-overlapping"
        )
    out = []
    for i in range(n_windows):
        out.append(
            window_trajectory(
                traj,
                start_sec=i * stride,
                duration_sec=duration_sec,
                window_id=i,
            )
        )
    return out


def default_mc_windows(
    *,
    sources: Iterable[str] = ("4T-M2-train",),
    duration_min: float = 10.0,
    n_windows: int = 3,
    start_min: float = 0.0,
) -> list[IntaSetTrajectory]:
    """Build the default pseudo-MC set for the real-data Pareto experiment.

    ``start_min`` skips an initial segment (e.g. data used for open-loop
    calibration) so evaluation windows do not overlap the fit set.
    """
    duration_sec = duration_min * 60.0
    start_sec = float(start_min) * 60.0
    windows: list[IntaSetTrajectory] = []
    for src in sources:
        traj = load_named(src)
        # Shift so windowing starts after ``start_min``.
        if start_sec > 0:
            m = traj.t_sec >= traj.t_sec[0] + start_sec - 1e-9
            if np.count_nonzero(m) < 2:
                raise ValueError(f"{src}: start_min={start_min} leaves no data")
            traj = IntaSetTrajectory(
                t_sec=traj.t_sec[m] - traj.t_sec[m][0],
                heights=traj.heights[m].copy(),
                pumps=traj.pumps[m].copy(),
                source=f"{traj.source}@t>={start_min:g}min",
            )
        # Shrink window length slightly if the holdout is just short of n*duration
        # (common when the recording ends at e.g. 2099.9 s).
        need = n_windows * duration_sec
        if traj.duration_sec + 1e-6 < need:
            duration_sec = traj.duration_sec / n_windows
        windows.extend(
            iter_windows(traj, duration_sec=duration_sec, n_windows=n_windows)
        )
    return windows
